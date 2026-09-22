# -*- coding: utf-8 -*-
"""Gera os slides da defesa (Defesa/slides_defesa.pptx) a partir das figuras da tese.

Porquê um script e não um .pptx feito à mão
As figuras da dissertação vivem em `Tese/images/resultados/` e já derivaram das
campanhas uma vez (21 jul: oito figuras desatualizadas no PDF). Um .pptx feito
à mão é uma segunda cópia que deriva em silêncio. Aqui cada slide aponta para a
figura da tese pelo nome; regenerar é correr o script. Os números escritos nos
slides são os da dissertação — o guião de defesa (docs/DEFESA_PERGUNTAS.md)
diz porquê: nunca defender um número que a tese não diga.

Cada slide leva NOTAS DO ORADOR com o que dizer, para servir de guião e de
material de estudo (docs/RESUMO_PARA_DECORAR.md tem a versão longa).

Uso: .venv/Scripts/python.exe Defesa/gerar_slides.py
"""
from __future__ import annotations

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FIG = os.path.join(RAIZ, "Tese", "images", "resultados")
IMG = os.path.join(RAIZ, "Tese", "images")
SAIDA = os.path.join(RAIZ, "Defesa", "slides_defesa.pptx")

# ── Paleta ────────────────────────────────────────────────────────────────────
# É a do painel (dashboard/theme.py): monocromática noturna, com a COR reservada
# para o que tem significado científico — as séries GNN/PPO/SAC. O que se copia
# aqui são os valores do **Modo Defesa** do painel, não os do ecrã: um
# videoprojetor achata os pretos, e o #7d7d7d que se lê no portátil desaparece na
# parede. Daí o fundo chapado (sem o gradiente do painel, que a lâmpada suja), as
# bordas a #2e2e2e em vez de #1f1f1f e os cinzentos subidos.
FUNDO = RGBColor(0x05, 0x05, 0x05)      # fundo dos slides
SUPERFICIE = RGBColor(0x0E, 0x0E, 0x0E)  # cartões e tabelas
SUPERFICIE2 = RGBColor(0x16, 0x16, 0x16)  # linhas alternadas
BORDA = RGBColor(0x2E, 0x2E, 0x2E)
INK = RGBColor(0xE2, 0xE2, 0xE2)        # texto corrente
INK_FORTE = RGBColor(0xF5, 0xF5, 0xF5)  # títulos e destaques
MUTED = RGBColor(0x9A, 0x9A, 0x9A)      # legendas e rótulos
LINHA = BORDA
BRANCO_CHAPA = RGBColor(0xFF, 0xFF, 0xFF)  # placa por baixo das figuras

# Séries: as mesmas famílias das figuras da tese, subidas para lerem sobre preto.
GNN = RGBColor(0x2F, 0xB5, 0x83)
PPO = RGBColor(0xEF, 0x78, 0x50)
SAC = RGBColor(0x5F, 0x9F, 0xE4)
# O painel não tem cor de acento — o destaque é o branco. Mantém-se o nome para
# o conteúdo dos slides não ter de mudar.
ACENTO = INK_FORTE

# Tipos de letra. O painel usa Space Grotesk / Inter / JetBrains Mono, que vêm do
# Google Fonts e NÃO estão instaladas nem nesta máquina nem, presumivelmente, na
# da sala: pedi-las ao PowerPoint faz cair num tipo arbitrário, e um slide com
# tipo de letra trocado nota-se mais do que um slide sem identidade. Usam-se os
# equivalentes mais próximos que qualquer Windows tem de origem, na mesma
# divisão de papéis: títulos / texto / números.
TITULO_FT = "Segoe UI Semibold"
CORPO_FT = "Segoe UI"
MONO_FT = "Consolas"

W, H = Inches(13.333), Inches(7.5)
MARGEM = Inches(0.6)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BRANCO = prs.slide_layouts[6]
_n = 0

# Os slides de reserva (ver o fim do ficheiro) numeram-se A1, A2, … e não entram
# na conta dos 20: quem folheia tem de perceber, pelo rodapé, que saiu do fio da
# apresentação. `_ANEXO` liga-se uma vez, antes de os construir.
_ANEXO = False
_na = 0

# O orçamento de tempo, em segundos por slide, na ordem em que são construídos.
# Quinze minutos é o que a capa promete; isto soma 14:40 e deixa 20 s de folga —
# um plano que já usa os quinze minutos todos falha no primeiro tropeção. Não é
# decoração: cada slide leva nas notas a hora a que deve começar e a que deve
# sair, que é a única forma de ensaiar sem cronómetro na mão e de saber, a meio,
# se se está adiantado ou atrasado. Se um slide for cortado ou acrescentado, esta
# lista tem de acompanhar — o `assert` no fim do ficheiro trava se deixar de bater.
# O roteiro (slide 2) entrou a 21 set com 20 s, pagos por cinco segundos à capa,
# ao problema, ao protocolo e à QI1 — todos ANTES da QI6, para que os dois
# relógios que a cábula manda decorar (7:25 ao entrar nela, 8:40 ao sair) não mudem.
PLANO = [20, 20, 45, 40, 60, 50, 45, 65, 40, 60, 75, 50, 30, 60, 40, 40, 30, 35, 60, 15]


def _mmss(seg):
    return "%d:%02d" % (seg // 60, seg % 60)


def _novo_slide():
    """Um slide com o fundo pintado. O layout em branco do PowerPoint é BRANCO —
    literalmente —, por isso o fundo escuro tem de ser um retângulo a cobrir a
    página, e tem de ser a PRIMEIRA forma (a ordem de inserção é a ordem z)."""
    s = prs.slides.add_slide(BRANCO)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    r.fill.solid()
    r.fill.fore_color.rgb = FUNDO
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def _texto(slide, x, y, w, h, linhas, tamanho=18, cor=INK, negrito=False,
           alinhar=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP, espaco=6, fonte=None):
    """Caixa de texto. `linhas` = str ou lista de str/(str, dict) — dict com
    tamanho/cor/negrito/nivel/fonte. Uma linha que comece por «• » é um marcador.

    O marcador sai em DUAS runs: um «▸» a cinzento e o texto na cor pedida. Com
    uma run só, o bullet herdava o peso e a cor da frase, e numa lista de sete
    pontos a negrito eram sete manchas a competir com o título. A cinzento, a
    lista lê-se pela primeira palavra de cada linha, que é o que se quer.
    """
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    if isinstance(linhas, str):
        linhas = [linhas]
    primeiro = True
    for item in linhas:
        txt, opc = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if primeiro else tf.add_paragraph()
        primeiro = False
        p.alignment = alinhar
        p.space_after = Pt(espaco)
        nivel = opc.get("nivel", 0)
        if nivel:
            p.level = nivel
        tam = opc.get("tamanho", tamanho)
        ft = opc.get("fonte", fonte) or CORPO_FT
        if txt.startswith("• "):
            marca = p.add_run()
            marca.text = "▸  "
            marca.font.size = Pt(tam)
            marca.font.color.rgb = MUTED
            marca.font.name = ft
            txt = txt[2:]
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(tam)
        r.font.bold = opc.get("negrito", negrito)
        r.font.color.rgb = opc.get("cor", cor)
        r.font.name = ft
    return tb


def _linha(slide, x1, y, x2, cor=BORDA, espessura=0.75):
    ln = slide.shapes.add_connector(1, x1, y, x2, y)
    ln.line.color.rgb = cor
    ln.line.width = Pt(espessura)
    return ln


def _e_numero(txt):
    """A célula é um valor numérico? Aceita o que estas tabelas usam à volta dos
    dígitos: vírgula decimal, %, ±, ≈, ×, /, parênteses, sinais e o «·» que
    separa as três medidas de uma célula da tabela de avaliação («38,3 · 86 % ·
    5/7») — que é tão coluna de números como as outras, e desalinhava do mesmo
    modo em tipo proporcional."""
    t = txt.strip()
    return bool(t) and any(c.isdigit() for c in t) and all(
        c.isdigit() or c in " ,.%±≈×/()+-—…·" for c in t)


def _borda_celula(cel, cor=BORDA, espessura=0.75):
    """Pinta as quatro arestas de uma célula.

    O python-pptx não expõe bordas de tabela: sem isto, o PowerPoint desenha as
    do estilo por omissão — claras e grossas, pensadas para tabelas sobre branco.
    Num slide preto ficam uma gaiola luminosa por cima do conteúdo. Escrevem-se
    à mão no XML da célula, e a ORDEM importa: o esquema exige
    lnL, lnR, lnT, lnB, e um elemento fora de ordem faz o PowerPoint declarar o
    ficheiro danificado.
    """
    tc = cel._tc
    pr = tc.get_or_add_tcPr()
    larg = str(int(espessura * 12700))
    hexa = "%02X%02X%02X" % (cor[0], cor[1], cor[2])
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for velho in pr.findall(qn(tag)):
            pr.remove(velho)
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = parse_xml(
            '<%s xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'w="%s" cap="flat" cmpd="sng" algn="ctr">'
            '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
            '<a:prstDash val="solid"/></%s>' % (tag, larg, hexa, tag))
        pr.append(ln)


def _rodape(slide):
    global _n, _na
    if _ANEXO:
        _na += 1
        etiqueta, legenda = "A%d" % _na, "Slide de reserva · não faz parte da apresentação"
    else:
        _n += 1
        # O número em Consolas, com zero à esquerda: alinha à direita sem dançar
        # entre o 9 e o 10, que numa numeração proporcional se nota ao folhear.
        etiqueta, legenda = "%02d" % _n, "Aprendizagem por Reforço Multiagente e Neuroevolução · ISCTE-IUL 2026"
    _linha(slide, MARGEM, H - Inches(0.55), W - MARGEM)
    _texto(slide, MARGEM, H - Inches(0.5), Inches(9), Inches(0.4), legenda,
           tamanho=10, cor=MUTED)
    _texto(slide, W - MARGEM - Inches(1), H - Inches(0.5), Inches(1), Inches(0.4),
           etiqueta, tamanho=10, cor=MUTED, alinhar=PP_ALIGN.RIGHT, fonte=MONO_FT)


def _titulo(slide, titulo, sub=""):
    _texto(slide, MARGEM, Inches(0.32), W - 2 * MARGEM, Inches(0.85), titulo,
           tamanho=30, negrito=True, cor=INK_FORTE, fonte=TITULO_FT)
    if sub:
        _texto(slide, MARGEM, Inches(1.02), W - 2 * MARGEM, Inches(0.5), sub,
               tamanho=15, cor=MUTED)
    # A régua do painel: um fio que atravessa a página e um segmento aceso à
    # esquerda. O painel faz isto com um gradiente no topo de cada cartão; aqui,
    # em que só há linhas retas, o mesmo efeito consegue-se com duas.
    _linha(slide, MARGEM, Inches(1.5), W - MARGEM)
    _linha(slide, MARGEM, Inches(1.5), MARGEM + Inches(1.1), cor=INK_FORTE, espessura=2.5)


_CACHE_FUNDO = {}


def _fundo_claro(caminho):
    """A figura tem fundo claro? Lê-se o canto superior esquerdo, que em todas as
    figuras aqui usadas é margem — o matplotlib deixa-a, e uma captura de ecrã
    começa pelo cromo da aplicação. Um pixel só chegava; usa-se um quadrado de
    8×8 para não decidir num pixel de anti-aliasing."""
    if caminho not in _CACHE_FUNDO:
        im = Image.open(caminho).convert("L").crop((0, 0, 8, 8))
        # `.resize((1, 1))` faz a média dos 64 pixels dentro do Pillow, e evita o
        # `getdata()`, que está a caminho de sair na versão 14.
        _CACHE_FUNDO[caminho] = im.resize((1, 1), Image.BOX).getpixel((0, 0)) > 128
    return _CACHE_FUNDO[caminho]


def _fig(slide, nome, x, y, w=None, h=None, pasta=FIG, chapa=True, folga=0.09):
    """Coloca a figura, por omissão sobre uma PLACA BRANCA.

    As figuras da tese têm fundo branco. Assentes diretamente no slide preto,
    ficam retângulos brancos a flutuar, com a margem interna do matplotlib a
    fazer de moldura irregular. A placa resolve-o ao contrário: o branco da
    figura funde-se com o branco da placa, e o que se vê é um cartão claro com
    aresta definida — o mesmo padrão que a cábula usa.
    """
    p = os.path.join(pasta, nome)
    if not os.path.exists(p):
        _texto(slide, x, y, w or Inches(4), Inches(0.6), "figura em falta: %s" % nome,
               tamanho=12, cor=PPO)
        return None
    if w is None or h is None:
        return slide.shapes.add_picture(p, x, y, width=w, height=h)

    # Cabe na caixa, sem deformar. A placa rouba a folga aos dois lados.
    pad = Inches(folga) if chapa else 0
    iw, ih = Image.open(p).size
    esc = min((w - 2 * pad) / iw, (h - 2 * pad) / ih)
    pw, ph = int(iw * esc), int(ih * esc)
    # Centrada na caixa pedida, para que uma figura estreita não fique encostada.
    px = x + (w - pw) // 2
    py = y + (h - ph) // 2
    if chapa:
        c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   px - pad, py - pad, pw + 2 * pad, ph + 2 * pad)
        c.fill.solid()
        # A placa segue o FUNDO da figura, não uma regra fixa: as figuras da tese
        # são brancas, mas a captura do painel é preta, e uma placa branca por
        # baixo dela punha uma auréola em volta de um screenshot escuro.
        c.fill.fore_color.rgb = BRANCO_CHAPA if _fundo_claro(p) else SUPERFICIE
        c.line.color.rgb = BORDA
        c.line.width = Pt(0.75)
        c.shadow.inherit = False
        # O raio do PowerPoint é relativo ao lado MENOR: sem isto, uma placa
        # baixa e larga sai com cantos de cápsula.
        c.adjustments[0] = min(0.06, 90000.0 / max(pw + 2 * pad, 1))
    return slide.shapes.add_picture(p, px, py, width=pw, height=ph)


def _notas(slide, texto, termos=()):
    """O guião do orador, com a marca de tempo à cabeça e, no fim, a CÁBULA dos
    termos que o slide usa — uma linha por termo, para o que se sabe mas não se
    consegue dizer de repente. Fica depois do guião, e separada, porque não é
    para dizer: é para consultar se a pergunta for «o que quer dizer X?».

    A marca vem do PLANO, não é escrita à mão: um slide que mude de sítio leva o
    seu tempo consigo, e dezanove horas copiadas para dentro de dezanove blocos
    de texto ficavam erradas à primeira troca. O índice é a posição do slide na
    apresentação, e não o número do rodapé — que salta a capa e o «Obrigado», por
    não os numerar, e por isso fica sempre uma unidade atrás.
    """
    if _ANEXO:
        marca = "[reserva — fora do tempo da apresentação; só se a pergunta vier]"
    else:
        i = len(prs.slides) - 1
        ini = sum(PLANO[:i])
        marca = "[%s → %s · %d s]" % (_mmss(ini), _mmss(ini + PLANO[i]), PLANO[i])
    corpo = marca + chr(10) * 2 + texto.strip()
    if termos:
        corpo += (chr(10) * 2 + "- - - CÁBULA: os termos deste slide (não é para dizer) - - -" + chr(10) * 2
                  + chr(10).join("• %s — %s" % t for t in termos))
    slide.notes_slide.notes_text_frame.text = corpo


def slide_texto_figura(titulo, sub, bullets, figura=None, notas="", fig_w=6.2, fig_h=4.9,
                       largura_texto=None, figuras=None, tamanho=17, termos=()):
    """Texto à esquerda, uma figura (ou várias empilhadas) à direita."""
    s = _novo_slide()
    _titulo(s, titulo, sub)
    lt = Inches(largura_texto) if largura_texto else (W - 2 * MARGEM - Inches(fig_w) - Inches(0.4)
                                                       if (figura or figuras) else W - 2 * MARGEM)
    _texto(s, MARGEM, Inches(1.75), lt, H - Inches(2.6), bullets, tamanho=tamanho, espaco=8)
    if figura:
        _fig(s, figura, W - MARGEM - Inches(fig_w), Inches(1.75), Inches(fig_w), Inches(fig_h))
    if figuras:
        n = len(figuras)
        alt = (Inches(fig_h) - Inches(0.2) * (n - 1)) / n
        y = Inches(1.75)
        for f in figuras:
            _fig(s, f, W - MARGEM - Inches(fig_w), y, Inches(fig_w), alt)
            y += alt + Inches(0.2)
    _rodape(s)
    _notas(s, notas, termos)
    return s


def slide_figura(titulo, sub, figura, legenda="", notas="", pasta=FIG, termos=()):
    s = _novo_slide()
    _titulo(s, titulo, sub)
    _fig(s, figura, MARGEM, Inches(1.7), W - 2 * MARGEM, H - Inches(2.9), pasta=pasta)
    if legenda:
        _texto(s, MARGEM, H - Inches(1.15), W - 2 * MARGEM, Inches(0.5), legenda,
               tamanho=13, cor=MUTED, alinhar=PP_ALIGN.CENTER)
    _rodape(s)
    _notas(s, notas, termos)
    return s


def slide_tabela(titulo, sub, cabecalho, linhas, notas="", larguras=None, bullets=None,
                 tamanho=13, termos=()):
    s = _novo_slide()
    _titulo(s, titulo, sub)
    y = Inches(1.75)
    if bullets:
        # A altura reservada aos marcadores era fixa (1,2"), e três frases longas
        # a 16 pt ocupam quatro linhas: na QI2 o terceiro marcador entrava por
        # cima do cabeçalho da tabela. Estima-se o número de linhas pelo
        # comprimento — a 16 pt, numa caixa de 12,1", cabem ~105 caracteres — e a
        # tabela desce o que for preciso. É uma estimativa, não uma medição: o
        # PowerPoint só sabe as quebras reais quando renderiza. Fica generosa de
        # propósito, porque o custo de sobrar espaço é nenhum e o de faltar é um
        # slide ilegível.
        linhas_est = sum(max(1, -(-len(t if isinstance(t, str) else t[0]) // 105))
                         for t in bullets)
        _texto(s, MARGEM, y, W - 2 * MARGEM, Inches(0.3) * linhas_est, bullets,
               tamanho=16, espaco=6)
        y += Inches(0.32) * linhas_est + Inches(0.25)
    n_l, n_c = len(linhas) + 1, len(cabecalho)
    altura = min(Inches(0.42) * n_l, H - y - Inches(0.8))
    tb = s.shapes.add_table(n_l, n_c, MARGEM, y, W - 2 * MARGEM, altura).table
    if larguras:
        for i, lw in enumerate(larguras):
            tb.columns[i].width = Inches(lw)
    # O PowerPoint aplica um estilo de tabela às riscas azuis por omissão, com
    # cabeçalho a cheio: sobre um slide preto é a única coisa que se vê. Desligam-se
    # as bandas e pinta-se célula a célula, como no painel — cabeçalho em
    # maiúsculas pequenas e cinzentas, corpo em duas superfícies alternadas.
    tb.first_row = False
    tb.horz_banding = False
    for j, c in enumerate(cabecalho):
        cel = tb.cell(0, j)
        cel.text = c.upper()
        r = cel.text_frame.paragraphs[0].font
        r.size, r.bold, r.color.rgb = Pt(tamanho - 2), True, MUTED
        r.name = MONO_FT
        cel.fill.solid()
        cel.fill.fore_color.rgb = FUNDO
        _borda_celula(cel)
    for i, linha in enumerate(linhas, start=1):
        for j, v in enumerate(linha):
            cel = tb.cell(i, j)
            cel.text = str(v)
            r = cel.text_frame.paragraphs[0].font
            r.size = Pt(tamanho)
            # A primeira coluna é o rótulo da linha: é ela que se lê a saltar, e
            # por isso vai mais clara do que os valores.
            r.color.rgb = INK_FORTE if j == 0 else INK
            r.bold = (j == 0)
            # Números em Consolas, como no painel: uma coluna de «3,74 / 3,29 /
            # 1,95» em tipo proporcional não alinha nas vírgulas, e é a coluna
            # que se lê de relance. Só quando a célula É um número — «12,8» sim,
            # «MARL por gradiente» não.
            r.name = MONO_FT if _e_numero(str(v)) else CORPO_FT
            cel.fill.solid()
            cel.fill.fore_color.rgb = SUPERFICIE2 if i % 2 else SUPERFICIE
            _borda_celula(cel)
    _rodape(s)
    _notas(s, notas, termos)
    return s


def _logo_branco(nome):
    """Devolve uma cópia do logótipo com o traço a branco, para o fundo escuro.

    O `iscte.png` é arte preta sobre transparente: assente num slide preto,
    desaparece. Recolorem-se os pixels opacos para branco e preserva-se o canal
    alfa — a versão monocromática que as normas de identidade preveem para
    fundos escuros. Fica em `Defesa/.assets/`, gerado, não versionado à mão.
    """
    orig = os.path.join(IMG, nome)
    if not os.path.exists(orig):
        return None
    pasta = os.path.join(RAIZ, "Defesa", ".assets")
    os.makedirs(pasta, exist_ok=True)
    dest = os.path.join(pasta, nome.replace(".png", "_branco.png"))
    im = Image.open(orig).convert("RGBA")
    alfa = im.getchannel("A")
    branco = Image.new("RGBA", im.size, (255, 255, 255, 255))
    branco.putalpha(alfa)
    branco.save(dest)
    return dest


# ── Cábula de termos, slide a slide ───────────────────────────────────────────
# Vai para o fim das notas de cada slide (ver `_notas`). As definições são as do
# vocabulário do docs/RESUMO_PARA_DECORAR.md (secção 11) e da tese, encurtadas a
# uma linha: o que se consegue ler de relance no ecrã do orador.
TERMOS = {
    "problema": [
        ("Robótica de enxame", "muitos robôs simples, sem líder; o comportamento coletivo nasce das interações locais"),
        ("Controlo descentralizado", "cada robô decide só com o que ele próprio observa — não há computador central"),
        ("MARL", "aprendizagem por reforço multiagente: vários agentes aprendem por tentativa e erro a maximizar uma recompensa"),
        ("Otimização bio-inspirada", "aqui, neuroevolução: evoluem-se os pesos da rede por seleção e mutação, sem gradientes"),
        ("Não-estacionaridade", "para cada agente o ambiente muda enquanto aprende, porque os outros agentes também estão a aprender"),
        ("Recompensa esparsa", "o sinal útil chega raramente (só ao chegar ao ninho): difícil saber que ação o causou"),
        ("Benchmark", "comparação sistemática, com o mesmo simulador, cenários e protocolo para todos"),
    ],
    "questoes": [
        ("Zero-shot", "aplicar a política treinada com N = 20 a outro N sem qualquer retreino"),
        ("Fitness", "a nota que a evolução dá a cada genoma; é o que decide quem se reproduz"),
        ("Shaping", "termos extra na recompensa/fitness que guiam a aprendizagem antes de haver sucesso"),
        ("Homing", "o termo que premeia terminar perto do ninho"),
        ("Deceção (deceptive)", "cenário em que seguir o gradiente da recompensa afasta da solução"),
        ("Procura por novidade", "premiar comportamentos diferentes dos já vistos, e não só o objetivo"),
    ],
    "simulador": [
        ("Ninho", "a zona-alvo; conta uma chegada quando os agentes exigidos estão lá ao mesmo tempo"),
        ("LiDAR", "sensor de distância por raios: diz a que distância está a parede em cada direção (alcance 8 m)"),
        ("Observação local e parcial", "cada robô só vê à sua volta e em relação a si (egocêntrica), nunca o mapa inteiro"),
        ("Muro em U", "beco virado para o ninho: ir a direito prende o enxame; é preciso afastar-se para contornar"),
        ("Porta Cooperativa", "a porta só abre com três robôs em simultâneo junto dela"),
        ("Perceção Cooperativa", "não tem ninho: conta uma chegada quando três agentes cercam um alvo móvel"),
        ("Porta com Alternativa", "a porta parece o caminho mais curto mas está fechada; há um corredor alternativo"),
        ("Wall-sliding", "contra um muro o robô não para — desliza ao longo da parede"),
    ],
    "controladores": [
        ("PPO", "Proximal Policy Optimization: gradiente de política on-policy, com passos cortados para não mudar demais"),
        ("SAC", "Soft Actor-Critic: off-policy, reutiliza experiência guardada e premeia a entropia (temperatura α)"),
        ("On-policy / off-policy", "treina só com dados da política atual / reaproveita dados antigos (replay buffer)"),
        ("RS2C", "Robust and Scalable Swarm Control: o framework da tese para PPO/SAC — execução descentralizada com partilha de parâmetros"),
        ("Parameter sharing", "os 20 agentes usam a mesma rede (os mesmos pesos)"),
        ("MLP", "rede neuronal densa clássica; entrada de tamanho fixo (aqui um vetor de 111 números, ℝ¹¹¹)"),
        ("GNN com atenção", "rede sobre o grafo dos vizinhos que pesa cada vizinho pela relevância; funciona com qualquer nº de vizinhos"),
        ("Invariante a N", "a mesma rede serve para 10, 20 ou 100 agentes sem mudar nada"),
        ("Shaping geodésico", "recompensa por aproximar-se do ninho, medida pelo caminho que contorna paredes"),
        ("Genoma / população", "genoma = um conjunto de pesos da rede; população = os 30 genomas avaliados por geração"),
        ("Ambientes vetorizados", "16 cópias do simulador a correr em paralelo para recolher experiência mais depressa"),
        ("Núcleos-hora", "núcleos × horas: a moeda que torna comparável o custo dos dois paradigmas"),
    ],
    "protocolo": [
        ("Execução (run)", "um treino completo, com a sua semente; é a unidade estatística (n = 7 por célula)"),
        ("Episódio", "uma tentativa do enxame, do início ao fim (2000 passos)"),
        ("Avaliação determinística", "o modelo já treinado joga sem aleatoriedade na ação, com sementes fixas"),
        ("Emparelhados / sementes comuns", "os três algoritmos enfrentam exatamente os mesmos 20 episódios"),
        ("Mann-Whitney U", "teste não-paramétrico: diz se um grupo tende a ter valores maiores que o outro, sem assumir normalidade"),
        ("δ de Cliff", "tamanho de efeito em [−1, 1]: P(A > B) − P(B > A); |δ| ≥ 0,474 é grande, ±1 é separação total"),
        ("Taxa de sucesso (P_task)", "fração de episódios com pelo menos uma chegada"),
        ("Pré-registo", "hipótese, testes e regra de decisão escritos ANTES de haver dados"),
        ("Hook de pre-commit", "script que corre sozinho antes de cada commit e o recusa se um número deixar de bater"),
    ],
    "qi1": [
        ("Célula", "uma combinação algoritmo × cenário (3 × 7 = 21 células)"),
        ("δ ≥ +0,71", "em ≥ 85 % dos pares de execuções o GNN supera o outro — efeito grande"),
        ("Bimodal", "as execuções dividem-se em dois grupos, resolve ou fica a zero, sem nada no meio"),
        ("3/7, 4/7, 2/7", "quantas das 7 execuções chegam a 100 % de sucesso"),
    ],
    "distribuicao": [
        ("Boxplot", "caixa com mediana e quartis; com 7 pontos, os quartis dependem de 1 ou 2 execuções"),
        ("Dot plot", "um ponto por execução: mostra a forma real da distribuição"),
        ("Gradiente geodésico", "a descida da distância-ao-ninho medida contornando paredes; em campo aberto não canaliza nada"),
        ("Degenerar", "a execução converge para um comportamento inútil (< 1 chegada/ep)"),
    ],
    "qi5": [
        ("Fitness exploitation", "a população maximiza o termo de shaping sem cumprir a tarefa — o reward hacking da evolução"),
        ("Farmável", "um termo que se acumula só por andar às voltas, sem nunca chegar ao ninho"),
        ("Planalto / pressão seletiva", "todos os genomas têm fitness parecida, e a seleção deixa de distinguir bons de maus"),
        ("Homing terminal", "clip((Φ₀ − Φ_T)/Φ₀, 0, 1): quanto se aproximou do ninho entre o INÍCIO e o FIM — vaguear não o aumenta"),
        ("Potencial geodésico Φ", "distância ao ninho pelo caminho mais curto que contorna paredes (Dijkstra numa grelha de 0,4 m)"),
        ("Euclidiano", "distância em linha reta: atravessa as paredes, por isso puxa o enxame para dentro do beco"),
        ("Atribuição de crédito", "saber que comportamento mereceu a nota; o homing resolve-o, não resolve descobrir o desvio"),
    ],
    "qi6": [
        ("Novelty search", "parte da seleção passa a premiar comportamento diferente; o campeão reportado é sempre o melhor no objetivo"),
        ("w", "o peso da novidade na seleção (w = 0,5: metade objetivo, metade novidade)"),
        ("Dosagem adaptativa", "w decai depois de a descoberta se sustentar: paga-se exploração só enquanto compra descoberta"),
        ("Orçamento igualado", "todos os braços com o mesmo tempo de treino (195 min), para a diferença não ser tempo"),
        ("Braço", "uma das condições comparadas (ex.: GNN adaptativo, GNN objetivo, PPO, SAC)"),
        ("Fisher exato", "teste para contagens (28/28 vs 15/28): quão improvável seria a diferença por acaso"),
        ("p = 0,088", "acima de 0,05: não há evidência de diferença — «indistinguíveis»"),
    ],
    "qi2": [
        ("Zero-shot", "treinada com N = 20, testada com 10, 50 e 100 sem retreino"),
        ("Chegadas por agente", "chegadas por episódio a dividir por N — permite comparar enxames de tamanhos diferentes"),
        ("Retenção per capita", "chegadas por agente a N = 100 ÷ as de N = 20; mede diluição do recurso, não perda de coordenação"),
        ("ℝ¹¹¹", "a MLP recebe um vetor de exatamente 111 números; com outro N o vetor muda de tamanho e a rede não o aceita"),
    ],
    "qi3": [
        ("Agente inerte", "10 % dos robôs param de repente a meio do episódio e ficam como obstáculos"),
        ("Retenção de chegadas", "chegadas com falhas ÷ chegadas sem falhas (92–106 % nos três)"),
        ("Parameter sharing", "todos têm a mesma política: qualquer robô substitui o que falhou"),
    ],
    "qi7": [
        ("Mapa composto", "labirinto de 103 × 62 m com quatro dificuldades em série: gargalo+U, quatro salas, porta cooperativa, ninho"),
        ("Transferência sem retreino", "levar os modelos treinados nos cenários pequenos para o mapa grande tal como estão"),
        ("Condições de controlo", "variações que excluem causas triviais do zero (escala da observação, obstáculos, features da porta)"),
        ("Navegador geodésico", "controlador escrito à mão que segue o caminho mais curto; não aprende — prova que o mapa é resolúvel"),
        ("Treino nativo", "treinar de raiz no próprio mapa composto (21 execuções, 780 min cada)"),
        ("Limiar pré-registado", "15 execuções convergentes em 21, fixado antes dos dados; ficou em 4"),
    ],
    "qi4": [
        ("Mapa de escolha", "em vez de um vencedor: que método usar consoante a missão"),
        ("Dimensão fixa / variável", "sabe-se ou não quantos robôs vai ter o enxame na missão"),
        ("Núcleos-hora", "12,8 (PPO/SAC) contra 97,6 (evolutivo) por execução — ≈ 8×"),
    ],
    "limitacoes": [
        ("Arquitetura assimétrica", "o evolutivo usa GNN com atenção e o PPO/SAC uma MLP: não se separa otimizador de representação"),
        ("Limite inferior", "o SAC podia fazer melhor com mais treino/afinação; o valor medido é um mínimo"),
        ("Temperatura α", "no SAC, o peso da entropia (exploração); aqui fixa em 0,1, sem o ajuste automático («dual»)"),
        ("Deployment gap", "o fosso entre simulação e robô físico — não validado nesta tese"),
        ("Costuras da física", "dois sítios onde um robô consegue atravessar a parede (no teto e nas junções em T)"),
        ("Célula contaminada", "resultado afetado pelas costuras: declarado, e não usado como evidência"),
    ],
    "contributos": [
        ("Contributo metodológico", "sobre como fazer (o desenho da fitness), não sobre qual algoritmo é melhor"),
        ("Verificador", "script que recalcula um número da tese a partir dos CSV e falha se não bater"),
    ],
    "demo": [
        ("Proveniência", "no painel, de onde vem cada número: o ficheiro CSV e a linha que o produziu"),
        ("Episódio 3D", "a vista do painel que desenha os enxames em 3D, a correr ao mesmo tempo"),
    ],
}

# 1. Capa
s = _novo_slide()
# Uma barra de acento à esquerda do título, do painel: o destaque é a luz, não
# uma cor institucional.
_barra = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGEM, Inches(2.15), Pt(3), Inches(2.05))
_barra.fill.solid()
_barra.fill.fore_color.rgb = INK_FORTE
_barra.line.fill.background()
_barra.shadow.inherit = False
_CAPA_X = MARGEM + Inches(0.32)
_texto(s, _CAPA_X, Inches(1.55), W - _CAPA_X - MARGEM, Inches(0.4),
       "DISSERTAÇÃO DE MESTRADO · DEFESA", tamanho=12, cor=MUTED, fonte=MONO_FT)
_texto(s, _CAPA_X, Inches(2.05), W - _CAPA_X - MARGEM - Inches(1.0), Inches(1.6),
       "Aprendizagem por Reforço Multiagente e Neuroevolução no Controlo de Enxames de Robôs:",
       tamanho=28, negrito=True, cor=INK_FORTE, fonte=TITULO_FT)
_texto(s, _CAPA_X, Inches(3.45), Inches(9.4), Inches(1.0),
       "Desempenho, Escalabilidade e Robustez", tamanho=22, cor=MUTED)
_linha(s, MARGEM, Inches(4.75), W - MARGEM)
_texto(s, MARGEM, Inches(5.0), W - 2 * MARGEM, Inches(1.2), [
    ("Gonçalo Pombo", {"tamanho": 20, "negrito": True, "cor": INK_FORTE}),
    ("Orientador: Prof. Doutor Luís Nunes", {"tamanho": 16, "cor": MUTED}),
    ("Mestrado em Inteligência Artificial · ISCTE-IUL · 2026", {"tamanho": 16, "cor": MUTED}),
])
_logo = _logo_branco("iscte.png")
if _logo:
    _iw, _ih = Image.open(_logo).size
    _lw = Inches(1.9)
    s.shapes.add_picture(_logo, W - MARGEM - _lw, Inches(0.55),
                         width=_lw, height=int(_lw * _ih / _iw))
_notas(s, """
Bom dia. Vou apresentar a dissertação «Aprendizagem por Reforço Multiagente e Neuroevolução no Controlo de Enxames de Robôs: Desempenho,
Escalabilidade e Robustez»: uma comparação, no mesmo simulador e com o mesmo protocolo, entre dois
paradigmas de controlo descentralizado — aprendizagem por reforço multiagente por
gradiente (PPO e SAC) e neuroevolução de uma rede de grafos com atenção — em oito
cenários. Quinze minutos: o problema, o método, sete perguntas e as respostas.

- - - NÃO É PARA DIZER: o Resumo da dissertação, que o júri leu (235 palavras) - - -

Esta dissertação compara a aprendizagem por reforço multiagente e a otimização
bio-inspirada no controlo descentralizado de enxames de robôs. Num simulador
próprio, vinte robôs têm de chegar em conjunto a um ninho. Um controlador
evolutivo, uma rede neuronal sobre grafos com atenção, foi comparado com dois
algoritmos de aprendizagem por reforço por gradiente (Proximal Policy
Optimization e Soft Actor-Critic) em sete cenários de dificuldade crescente, com
7 treinos independentes por algoritmo e cenário.

Com uma função de aptidão que premeia também a aproximação ao ninho, o
controlador evolutivo supera os métodos de gradiente em três cenários, iguala o
melhor noutros dois e é o único que funciona sem novo treino com enxames de 10 a
100 robôs; os métodos de gradiente são mais fiáveis em espaço aberto e gastam
menos cálculo. O cenário que obriga a afastar-se do ninho para o alcançar só foi
resolvido de forma consistente com pressão por comportamentos novos, atenuada
após a descoberta: numa replicação com 28 treinos por algoritmo, resolveu-o em
28, contra 15 sem essa pressão e 14 de cada método de gradiente. Um oitavo
cenário, que junta as dificuldades dos sete num labirinto de 103 x 62 m, é
resolvido em apenas 4 de 21 treinos.

A hipótese de que a aprendizagem adaptativa supera a otimização estática
confirma-se só em parte: escalar depende da representação da política, não do
otimizador; nos cenários enganadores, depende do sinal de treino.

É daqui que saem, quase sempre, as primeiras perguntas. Cada afirmação deste
Resumo tem o seu slide: a fitness no 10, a novidade no 11, a escala no 12, o mapa
composto no 14, e a frase final é o slide 18.
""")

# 2. Roteiro — o índice da apresentação
# Cinco partes, cada uma com os rótulos de rodapé que ocupa. Com o roteiro a
# levar o 01, o rótulo de cada slide passa a ser a sua posição menos um — a capa
# continua sem número —, e é por esses rótulos que o júri pode pedir «volte ao 10».
s = _novo_slide()
_titulo(s, "Roteiro", "cinco partes, quinze minutos")
_PARTES = [
    ("01", "Contexto", "o problema, a lacuna na literatura e as sete questões de investigação", "02–03"),
    ("02", "Método", "o simulador e os oito cenários, os três controladores, o protocolo experimental", "04–06"),
    ("03", "Resultados", "as sete questões, pela ordem em que se desbloqueiam: QI1, 5, 6, 2, 3, 7 e 4",
     "07–14"),
    ("04", "Balanço", "as limitações, os contributos e a conclusão", "15–17"),
    ("05", "Demo ao vivo", "o Muro em U no painel, com a proveniência de cada número", "18"),
]
_ry, _rh = Inches(1.8), Inches(0.9)
for _k, (_num, _nome, _desc, _rot) in enumerate(_PARTES):
    _y = _ry + _k * _rh
    _texto(s, MARGEM, _y + Inches(0.08), Inches(1.0), Inches(0.6), _num,
           tamanho=28, cor=MUTED, fonte=MONO_FT)
    _texto(s, MARGEM + Inches(1.1), _y + Inches(0.05), Inches(8.9), Inches(0.4), _nome,
           tamanho=21, negrito=True, cor=INK_FORTE, fonte=TITULO_FT)
    _texto(s, MARGEM + Inches(1.1), _y + Inches(0.45), Inches(8.9), Inches(0.4), _desc,
           tamanho=15, cor=MUTED)
    _texto(s, W - MARGEM - Inches(2.0), _y + Inches(0.18), Inches(2.0), Inches(0.4),
           ("slides " if "–" in _rot else "slide ") + _rot, tamanho=14, cor=MUTED, fonte=MONO_FT, alinhar=PP_ALIGN.RIGHT)
    if _k < len(_PARTES) - 1:
        _linha(s, MARGEM, _y + _rh - Inches(0.03), W - MARGEM)
_rodape(s)
_notas(s, """
Em cinco partes. Primeiro o contexto: o problema e as sete questões. Depois o
método — simulador, controladores e protocolo. O grosso é o dos resultados, que
respondo pela ordem em que se desbloqueiam, e não pela ordem de numeração. Fecho
com as limitações, os contributos e a conclusão, e um minuto de demo ao vivo.

(Os números à direita são os rótulos do rodapé: se o júri pedir «volte à QI5»,
é o 09.)
""")

# 3. Problema
slide_texto_figura(
    "O problema", "duas escolas, e nenhuma comparação direta em cenários difíceis",
    [
        "• Robótica de enxame: muitos agentes simples, controlo descentralizado, "
        "inteligência nas interações locais",
        "• Duas escolas para desenhar o controlador:",
        ("Otimização bio-inspirada — robustez offline, controladores estáticos", {"nivel": 1, "tamanho": 15}),
        ("MARL — adaptabilidade online, não-estacionaridade, recompensas esparsas", {"nivel": 1, "tamanho": 15}),
        "• Lacuna (Majid et al., 2024; Bettini et al., 2024): faltam benchmarks diretos, "
        "com estatística, em cenários dinâmicos e difíceis",
        "• Hipótese de partida: a inteligência adaptativa (MARL) supera a robustez "
        "estática bio-inspirada em cenários de stress?",
        ("A resposta curta: só em parte — e não onde a hipótese a punha.", {"negrito": True, "cor": ACENTO}),
    ],
    figura="viz_u_wall.png", fig_w=5.6, fig_h=4.6,
    termos=TERMOS["problema"],
    notas="""
O campo organiza-se em duas escolas. A bio-inspirada trata o controlador como um
problema de otimização offline — robusto, mas estático. O MARL aprende online,
mas paga a não-estacionaridade e as recompensas esparsas. A literatura diz que
faltam comparações diretas, rigorosas, em cenários difíceis. Esta tese é essa
comparação. A hipótese de partida era que a adaptabilidade do MARL ganharia em
stress; a resposta vai ser «só em parte» — e a parte que se confirma não está
onde a hipótese a punha.
""")

# 3. Questões
slide_texto_figura(
    "Sete questões de investigação", "quatro de comparação entre paradigmas, três sobre os mecanismos",
    [
        ("Nível 1 — comparação", {"negrito": True, "cor": ACENTO}),
        "• QI1  Desempenho de tarefa: qual paradigma é mais eficaz, e em que cenários?",
        "• QI2  Escalabilidade: uma política treinada com N=20 transfere-se para N=10…100 sem retreino?",
        "• QI3  Robustez: o que acontece quando 10 % dos agentes falham a meio?",
        "• QI4  Critério de escolha: quando preferir um ao outro?",
        ("Nível 2 — mecanismos", {"negrito": True, "cor": ACENTO}),
        "• QI5  Desenho da fitness: um shaping de homing desbloqueia o que a fitness pura não resolve?",
        "• QI6  Deceção: a procura por novidade ajuda onde o gradiente engana?",
        "• QI7  Composição: as conclusões transferem-se para um mapa que junta as dificuldades?",
    ],
    termos=TERMOS["questoes"],
    notas="""
Sete questões em dois níveis. As quatro primeiras comparam paradigmas:
desempenho, escalabilidade, robustez e critério de escolha. As três seguintes
descem aos mecanismos — o sinal de treino, a pressão de exploração e a
topologia — e só se tornam formuláveis depois das primeiras. Vou respondê-las
todas, pela ordem em que os resultados as desbloqueiam.
""")

# 4. Simulador e cenários
s = _novo_slide()
_titulo(s, "O simulador e os oito cenários",
        "navegação cooperativa até ao ninho em 3D, 20 agentes, observação local (LiDAR 8 m)")
_texto(s, MARGEM, Inches(1.75), Inches(5.2), Inches(4.8), [
    "• Tarefa: chegar ao ninho em conjunto, sem objetos para apanhar; perceção cooperativa com alvo móvel",
    "• Observação local e parcial: LiDAR, bússola ao ninho, vizinhos",
    "• Sete cenários de dificuldade isolada:",
    ("Sandbox · Muro em U (deceção espacial) · Gargalo · Quatro Salas", {"nivel": 1, "tamanho": 14}),
    ("Porta Cooperativa (3 robôs) · Perceção Cooperativa · Porta com Alternativa", {"nivel": 1, "tamanho": 14}),
    "• Oitavo: o mapa composto, 103 × 62 m, que junta quatro dificuldades em série",
    "• Falhas súbitas de agentes e dimensão do enxame variável, para a robustez e a escala",
], tamanho=16, espaco=8)
grelha = ["mapa_3d_none.png", "mapa_3d_u_wall.png", "mapa_3d_bottleneck.png", "mapa_3d_four_rooms.png",
          "mapa_3d_cooperative_door.png", "mapa_3d_cooperative_perception.png",
          "mapa_3d_cooperative_door_bypass.png"]
x0, y0, cw, ch = W - MARGEM - Inches(7.2), Inches(1.75), Inches(1.7), Inches(1.45)
for k, f in enumerate(grelha):
    _fig(s, f, x0 + (k % 4) * (cw + Inches(0.13)), y0 + (k // 4) * (ch + Inches(0.13)), cw, ch, pasta=IMG)
# O oitavo, a planta do mapa composto, ocupa a linha de baixo inteira.
_fig(s, "mapa_grande_planta.png", x0, y0 + 2 * (ch + Inches(0.13)), Inches(7.2), Inches(2.0))
_rodape(s)
_notas(s, """
O simulador foi construído de raiz em Python: tridimensional, com física de
colisões e deslizamento nos muros, LiDAR horizontal de 8 metros e observação
egocêntrica. A tarefa é chegar ao ninho em conjunto: não há objetos para apanhar, e
conta uma chegada quando os agentes exigidos estão no ninho ao mesmo tempo. Sete cenários isolam uma
dificuldade cada — um beco enganador, um gargalo, quatro salas, uma porta que só
abre com três robôs — e o oitavo compõe quatro delas num labirinto quatro vezes
maior. Tudo é reproduzível: sementes fixas, avaliação determinística.
""", TERMOS["simulador"])

# 5. Controladores
slide_tabela(
    "Os três controladores", "o mesmo espaço de observação e ação; execução descentralizada com partilha de parâmetros",
    ["", "PPO / SAC (framework RS2C)", "Neuroevolução (GNN)"],
    [
        ["Paradigma", "MARL por gradiente (on-policy / off-policy)", "Otimização bio-inspirada, sem gradientes"],
        ["Política", "MLP sobre observação de vizinhança densa (entrada fixa, ℝ¹¹¹)", "Rede de grafos com atenção sobre os vizinhos (invariante a N)"],
        ["Sinal de treino", "Recompensa por passo com shaping geodésico", "Fitness: chegadas + homing terminal no potencial geodésico"],
        ["Orçamento por execução", "48 min, 16 ambientes vetorizados", "195 min, população de 30 genomas em paralelo"],
        ["Custo em núcleos-hora", "12,8", "97,6  (≈ 8×)"],
    ],
    larguras=[2.4, 4.9, 4.9],
    bullets=[("A arquitetura difere entre paradigmas — é uma limitação declarada, e é o que permite atribuir "
              "a escalabilidade à representação e não ao otimizador.", {"cor": MUTED, "tamanho": 15})],
    termos=TERMOS["controladores"],
    notas="""
Os dois lados partilham o simulador, a observação e o protocolo de avaliação.
O PPO e o SAC usam as implementações de referência da Stable-Baselines3 — uma
MLP de entrada fixa. O controlador evolutivo otimiza os pesos de uma rede de
grafos com atenção sobre os vizinhos, invariante ao número de agentes. A
assimetria é declarada como primeira limitação: por isso mesmo, a conclusão
sobre escalabilidade é atribuída à representação, não ao algoritmo. O evolutivo
custa cerca de oito vezes mais núcleos-hora por execução.
""")

# 6. Protocolo
slide_texto_figura(
    "O protocolo", "o que impede esta tese de mentir",
    [
        "• 3 algoritmos × 7 cenários × 7 execuções independentes = 147 treinos",
        "• Avaliação determinística: 20 episódios emparelhados por modelo (sementes comuns) — "
        "2 940 episódios na campanha principal",
        "• A unidade estatística é a execução, não o episódio: Mann-Whitney U sobre as médias "
        "por execução (n = 7), δ de Cliff como tamanho de efeito",
        "• Métrica de tarefa pura: chegadas por episódio e taxa de sucesso — comparável entre paradigmas",
        "• Três campanhas pré-registadas (hipótese, testes e regra de decisão fixados antes dos dados)",
        "• 27 verificadores automáticos, 19 no hook de pre-commit: cada número da tese é recalculado "
        "a partir dos CSV, e o commit é recusado se algum deixar de bater",
    ],
    figura="comparacao_barras_geral.png", fig_w=5.8, fig_h=4.4,
    termos=TERMOS["protocolo"],
    notas="""
Cento e quarenta e sete treinos, cada modelo avaliado em vinte episódios
determinísticos com as mesmas sementes para os três algoritmos. A unidade
estatística é a execução: sete por célula, Mann-Whitney com delta de Cliff. As
campanhas decisivas foram pré-registadas — hipótese, testes e regra de decisão
escritos antes de haver dados. E cada número da tese é conferido por um
verificador contra os ficheiros de avaliação, a cada commit. Isto é o que me
permite dizer que os números são os dos dados.
""")

# 7. QI1
slide_texto_figura(
    "QI1 — Desempenho de tarefa", "não há paradigma dominante; há cenários de cada um",
    [
        "• 15 das 21 células a 100 % de sucesso em todas as execuções",
        ("GNN evolutivo — especialista em navegação estruturada:", {"negrito": True, "cor": GNN}),
        "  superior em 3 cenários (Quatro Salas 59,8 vs 33,6 · Porta Cooperativa · Perceção Cooperativa, "
        "δ ≥ +0,71); empata com o PPO no Gargalo (121,4 vs 123,2) e na Porta com Alternativa",
        ("PPO — o generalista fiável:", {"negrito": True, "cor": PPO}),
        "  100 % em seis cenários, a menor variância; ganha o Sandbox (71,5 vs 38,3 do GNN, δ = −1,00)",
        ("SAC — fiável em espaço aberto, frágil nos gargalos físicos:", {"negrito": True, "cor": SAC}),
        "  o único que falha execuções no Gargalo (5/7, 41,4 ± 36,8)",
        ("Muro em U: ninguém o resolve de forma fiável — bimodal nos três (3/7, 4/7, 2/7).",
         {"negrito": True}),
    ],
    figuras=["taxa_sucesso_por_cenario.png", "recolhas_por_cenario.png"], fig_w=5.6, fig_h=5.0, tamanho=15,
    termos=TERMOS["qi1"],
    notas="""
Primeira resposta: não há vencedor universal. Quinze das vinte e uma células
estão a cem por cento. O evolutivo, com a fitness de homing, é o especialista em
navegação estruturada — significativamente superior em três cenários, incluindo
os dois cooperativos, e empatado com o PPO nos gargalos. O PPO é o generalista:
converge em tudo menos no Muro em U, com a variância mais baixa, e ganha o
Sandbox onde o evolutivo degenera. O SAC é o mais sensível aos gargalos. E o
Muro em U não é resolvido de forma fiável por nenhum: cada execução ou aprende
o desvio ou não aprende nada.
""")

# 8. Fiabilidade / bimodalidade
slide_texto_figura(
    "A unidade é a execução — e a forma da distribuição conta", "porque se reportam pontos e não caixas",
    [
        "• Com n = 7, os quartis de um boxplot são ruído; a caixa cheia sugere densidade onde não há nenhuma",
        "• Muro em U: 4 execuções do GNN a zero e 3 entre 43 e 80 — bimodal, não «média 24,5»",
        "• Sandbox: o cenário mais simples é o menos fiável para o evolutivo — 2 execuções degeneram "
        "(sem paredes, não há gradiente geodésico que canalize o homing)",
        "• A convergência tudo-ou-nada é o padrão em 6 das 21 células — e é o que as QI5 e QI6 explicam",
    ],
    figuras=["dotplot_eval_u_wall.png", "dotplot_eval_none.png"], fig_w=5.4, fig_h=5.0,
    termos=TERMOS["distribuicao"],
    notas="""
Um ponto de método que muda a leitura. Com sete execuções, um boxplot esconde a
forma. No Muro em U, o GNN tem quatro execuções a zero e três a resolver o
cenário por completo: a média de vinte e quatro chegadas não descreve nenhuma
delas. No Sandbox, o cenário mais simples, o evolutivo é o menos fiável: sem
paredes não há gradiente geodésico que o guie, e duas execuções degeneram. Este
padrão tudo-ou-nada é o fio condutor das duas questões seguintes.
""")

# 9. QI5
s = slide_texto_figura(
    "QI5 — O desenho da fitness decide", "o «colapso do evolutivo» era um artefacto do sinal de treino",
    [
        "• Fitness inicial: chegadas + retorno acumulado do episódio — farmável por deambulação; a população "
        "satura num planalto sem pressão seletiva (fitness exploitation)",
        "• Cura: substituir o retorno pelo homing terminal — proximidade ao ninho no fim do episódio, medida "
        "no potencial geodésico (contorna paredes; o euclidiano atrai para dentro do beco)",
        ("Resultado: de 0 % para 100 % de sucesso nas 28 execuções dos quatro cenários de gargalo.",
         {"negrito": True, "cor": GNN}),
        "• Necessário, não suficiente: resolve a atribuição de crédito, não a descoberta do desvio comprido "
        "— o Muro em U continua bimodal (3/7), como nos métodos de gradiente",
    ],
    largura_texto=6.1, tamanho=16,
    termos=TERMOS["qi5"],
    notas="""
A primeira das questões de mecanismo. Nas campanhas exploratórias o evolutivo
colapsava nos labirintos, e a leitura fácil era «limitação do paradigma». Não
era. A fitness inicial somava o retorno acumulado, que se farma a deambular; a
população saturava. Substituí esse termo pelo homing terminal — estar perto do
ninho no fim, medido num potencial geodésico que contorna as paredes. Vinte e
oito execuções de zero para cem por cento. Mas é condição necessária, não
suficiente: o Muro em U continuou bimodal. Faltava a descoberta.
""")

# O quadro «antes → depois». Até 21 set este slide mostrava o mapa do potencial
# euclidiano contra o geodésico (heatmap_geodesico_u_wall.png): explica o
# mecanismo, mas com duas barras de cor, eixos em metros e um subtítulo em inglês
# não se lê projetado, e o que o slide afirma é o salto de 0 % para 100 %. Esse
# salto desenha-se aqui, cenário a cenário; o mecanismo fica para as notas.
_qx, _qy, _qw, _qh = W - MARGEM - Inches(5.6), Inches(1.75), Inches(5.6), Inches(4.85)
_q = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _qx, _qy, _qw, _qh)
_q.fill.solid()
_q.fill.fore_color.rgb = SUPERFICIE
_q.line.color.rgb = BORDA
_q.line.width = Pt(0.75)
_q.shadow.inherit = False
_q.adjustments[0] = 0.04
_px = _qx + Inches(0.35)
_texto(s, _px, _qy + Inches(0.25), _qw - Inches(0.7), Inches(0.4),
       "TAXA DE SUCESSO DO GNN EVOLUTIVO", tamanho=12, cor=MUTED, fonte=MONO_FT)
_c1, _c2 = _px + Inches(2.45), _px + Inches(3.75)
_texto(s, _c1, _qy + Inches(0.7), Inches(1.15), Inches(0.5), ["fitness", "inicial"], tamanho=11,
       cor=MUTED, fonte=MONO_FT, alinhar=PP_ALIGN.CENTER, espaco=0)
_texto(s, _c2, _qy + Inches(0.7), Inches(1.15), Inches(0.5), ["homing", "terminal"], tamanho=11,
       cor=GNN, fonte=MONO_FT, alinhar=PP_ALIGN.CENTER, espaco=0)
_ly = _qy + Inches(1.3)
for _k, _cen in enumerate(["Gargalo", "Quatro Salas", "Porta Cooperativa", "Porta com Alternativa"]):
    _y = _ly + _k * Inches(0.62)
    _linha(s, _px, _y, _qx + _qw - Inches(0.35))
    _texto(s, _px, _y + Inches(0.12), Inches(2.4), Inches(0.45), _cen, tamanho=15, cor=INK_FORTE,
           ancora=MSO_ANCHOR.MIDDLE)
    _texto(s, _c1, _y + Inches(0.08), Inches(1.15), Inches(0.5), "0 %", tamanho=22, cor=MUTED,
           fonte=MONO_FT, alinhar=PP_ALIGN.CENTER)
    _texto(s, _c1 + Inches(1.1), _y + Inches(0.1), Inches(0.3), Inches(0.5), "→", tamanho=18,
           cor=MUTED, alinhar=PP_ALIGN.CENTER)
    _texto(s, _c2, _y + Inches(0.08), Inches(1.15), Inches(0.5), "100 %", tamanho=22, negrito=True,
           cor=GNN, fonte=MONO_FT, alinhar=PP_ALIGN.CENTER)
_linha(s, _px, _ly + 4 * Inches(0.62), _qx + _qw - Inches(0.35))
_texto(s, _px, _ly + 4 * Inches(0.62) + Inches(0.2), _qw - Inches(0.7), Inches(0.9), [
    ("28 execuções, 7 por cenário: todas a 100 %", {"tamanho": 16, "negrito": True, "cor": INK_FORTE}),
    ("Muro em U, com o mesmo homing: 3/7 — bimodal", {"tamanho": 14, "cor": MUTED}),
], espaco=4)

# 10. QI6
slide_texto_figura(
    "QI6 — Deceção e procura por novidade", "a exploração tem preço; doseada, deixa de ter",
    [
        "• Novidade com peso fixo (w = 0,5), orçamento igualado: Muro em U passa de 3/7 para 7/7 "
        "(69,8 vs 24,5; p = 0,026, δ = +0,71) — mas custa na Porta com Alternativa (63,0 vs 86,7; δ = −1,00)",
        "• Dosagem adaptativa (pré-registada): w decai após a descoberta sustentada → mantém 7/7 no Muro em U, "
        "sem custo significativo em nenhum dos outros seis; com o dobro do orçamento, 88,7 ± 0,6 — "
        "o melhor resultado da dissertação",
        "• O ganho não se compra com orçamento: o objetivo puro com 390 min continua bimodal (4/7)",
        ("Replicação a n = 28: 28/28 execuções resolvidas, contra 15/28 do objetivo puro e 14/28 do PPO "
         "e do SAC (Fisher exato, p < 0,0001).", {"negrito": True, "cor": ACENTO}),
        "• GNN objetivo e PPO indistinguíveis (p = 0,088): o que separa os braços é a dosagem da exploração, "
        "não o paradigma",
    ],
    figura="megatreino_u_wall_4bracos.png", fig_w=5.6, fig_h=5.0, tamanho=15,
    termos=TERMOS["qi6"],
    notas="""
A hibridização com procura por novidade, com orçamento igualado, resolve o
Muro em U em sete de sete execuções — a única configuração que o faz. Mas com
peso fixo degrada o cenário onde o gradiente já bastava. A campanha
pré-registada de dosagem adaptativa — pagar novidade só enquanto ela compra
descoberta — removeu o custo e deu o melhor resultado da tese. E a replicação
com vinte e oito execuções por braço fecha a questão: vinte e oito em vinte e
oito, contra quinze do objetivo puro e catorze de cada método de gradiente. O
GNN objetivo e o PPO são indistinguíveis; o que decide é a dosagem da
exploração, não o paradigma.
""")

# 11. QI2
slide_tabela(
    "QI2 — Escalabilidade Zero-Shot", "só a arquitetura de grafo transfere de N = 10 a N = 100 sem retreino",
    ["Cenário", "N=10", "N=20", "N=50", "N=100", "Retenção per capita"],
    [
        ["Sandbox", "3,74", "3,29", "1,95", "1,27", "39 %"],
        ["Perceção Cooperativa", "1,64", "1,30", "0,83", "0,59", "45 %"],
        ["Gargalo", "6,92", "6,93", "5,96", "4,04", "58 %"],
        ["Quatro Salas", "3,60", "3,86", "3,26", "2,54", "66 %"],
        ["Muro em U", "3,84", "4,01", "3,81", "3,14", "78 %"],
        ["Porta Cooperativa", "3,40", "3,54", "3,41", "3,12", "88 %"],
        ["Porta com Alternativa", "4,31", "4,50", "4,44", "4,06", "90 %"],
    ],
    larguras=[3.2, 1.6, 1.6, 1.6, 1.6, 2.5],
    bullets=[
        "• 100 % de sucesso nas 28 combinações cenário × dimensão (chegadas por agente na tabela)",
        "• PPO/SAC: a MLP de entrada fixa (ℝ¹¹¹) é estruturalmente incompatível com N ≠ 20 — não é um "
        "resultado fraco, é o resultado",
        ("• A retenção é maior nos cenários com paredes: a estrutura atenua a diluição do recurso, "
         "não a agrava. Zero-shot é uma propriedade da representação, não do otimizador.", {"negrito": True}),
    ],
    termos=TERMOS["qi2"],
    notas="""
Escalabilidade. A política do controlador de grafo, treinada com vinte agentes,
transfere para dez, cinquenta e cem — cem por cento de sucesso nas vinte e oito
combinações, incluindo labirintos e tarefas cooperativas, onde a objeção do
congestionamento seria mais plausível. A retenção per capita é mais alta nos
cenários com paredes do que nos abertos: o que se perde no Sandbox é diluição
de um recurso finito, não coordenação. O PPO e o SAC nem sequer podem ser
avaliados fora de N igual a vinte. A vantagem de escala está na representação.
""")

# 12. QI3
slide_texto_figura(
    "QI3 — Robustez a falhas", "10 % dos agentes ficam inertes a meio do episódio",
    [
        "• Retenção de chegadas entre 92 % e 106 % nas 21 combinações algoritmo–cenário",
        "• Inclui a Porta Cooperativa, que exige três agentes simultâneos: o enxame redistribui-se",
        "• Transversal aos paradigmas: a redundância vem do parameter sharing e da observação local — "
        "nenhum agente é insubstituível",
        "• Consequência para a QI4: a robustez não é critério de escolha entre paradigmas",
    ],
    figura="robustez_falhas.png", fig_w=6.6, fig_h=4.6,
    termos=TERMOS["qi3"],
    notas="""
Robustez: dez por cento dos agentes falham a meio do episódio e ficam inertes.
Os três paradigmas retêm entre noventa e dois e cento e seis por cento das
chegadas, em todas as vinte e uma combinações — mesmo na porta que precisa de
três robôs. A redundância vem da partilha de parâmetros e da observação local.
Sendo transversal, a robustez não discrimina entre paradigmas — e isso é uma
resposta, não uma ausência de resposta.
""")

# 13. QI7
# Até 21 set: a planta e os rastos empilhados numa coluna de 6" — a figura dos
# rastos é alta (três mapas, um por baixo do outro) e, encolhida a metade da
# altura do slide, os percursos deixavam de se ver. A planta já está no slide 4.
# Agora os três mapas dos rastos vão lado a lado, cada um recortado da figura da
# tese (`_paineis_rastos`), e o texto passa para baixo, mais curto.
def _paineis_rastos(nome="mapa_grande_rastos.png"):
    """Recorta os três mapas da figura dos rastos (GNN, PPO, SAC, de cima para
    baixo) para `Defesa/.assets/`. Os cortes medem-se na imagem — as faixas de
    linhas com conteúdo com mais de 300 px de altura são os mapas; os títulos e a
    legenda são faixas finas —, para sobreviverem a uma figura regenerada."""
    import numpy as np
    orig = os.path.join(FIG, nome)
    im = Image.open(orig).convert("RGB")
    a = np.asarray(im).astype(int).sum(axis=2) < 740
    linhas = a.any(axis=1)
    faixas, ini = [], None
    for y, v in enumerate(list(linhas) + [False]):
        if v and ini is None:
            ini = y
        elif not v and ini is not None:
            if y - ini > 300:
                faixas.append((ini, y))
            ini = None
    assert len(faixas) == 3, "esperava 3 mapas em %s, encontrei %d" % (nome, len(faixas))
    pasta = os.path.join(RAIZ, "Defesa", ".assets")
    os.makedirs(pasta, exist_ok=True)
    saida = []
    for k, (y0, y1) in enumerate(faixas):
        xs = np.where(a[y0:y1].any(axis=0))[0]
        dest = os.path.join(pasta, "rastos_%d.png" % k)
        im.crop((max(xs.min() - 10, 0), max(y0 - 10, 0), xs.max() + 10, y1 + 10)).save(dest)
        saida.append(dest)
    return saida


s = _novo_slide()
_titulo(s, "QI7 — Composição de dificuldades",
        "um labirinto de 103 × 62 m que junta quatro dificuldades em série")
_cw, _gap = (W - 2 * MARGEM - 2 * Inches(0.15)) / 3, Inches(0.15)
for _k, ((_nome, _cor), _f) in enumerate(zip(
        [("GNN evolutivo · 8 chegadas", GNN), ("PPO · 0 chegadas", PPO), ("SAC · 0 chegadas", SAC)],
        _paineis_rastos())):
    _x = MARGEM + _k * (_cw + _gap)
    _texto(s, _x, Inches(1.62), _cw, Inches(0.32), _nome, tamanho=14, negrito=True, cor=_cor,
           alinhar=PP_ALIGN.CENTER)
    _fig(s, os.path.basename(_f), _x, Inches(1.95), _cw, Inches(2.5), pasta=os.path.dirname(_f))
_texto(s, MARGEM, Inches(4.47), W - 2 * MARGEM, Inches(0.3),
       "um episódio de cada, mesmas sementes · o ninho à direita · pontos brancos: partida; cheios: onde acabaram",
       tamanho=11, cor=MUTED, alinhar=PP_ALIGN.CENTER)
_texto(s, MARGEM, Inches(4.85), W - 2 * MARGEM, Inches(2.0), [
    "• Sem retreino: zero chegadas em 84 de 84 células (1 680 episódios) — e o mapa é resolúvel: "
    "um navegador geodésico sem aprendizagem faz 53,0 chegadas/ep",
    "• Treino nativo, 21 execuções por algoritmo: só o evolutivo passa — em 4 de 21, abaixo do "
    "limiar de 15 fixado antes dos dados; em 19 das 21 a fitness ainda subia no fim",
    ("A resposta à QI7 é negativa: a composição degrada a fiabilidade, não a magnitude — "
     "a bimodalidade do Muro em U em ponto grande.", {"negrito": True, "cor": ACENTO}),
], tamanho=15, espaco=6)
_rodape(s)
_notas(s, """
A última pergunta testa a objeção mais natural a um benchmark por cenários
isolados. Sem retreino, nenhum controlador faz uma única chegada no mapa
composto — em oitenta e quatro células, com controlos que excluem a escala da
observação e os obstáculos como causa. E o mapa é resolúvel: um navegador
geodésico sem aprendizagem faz cinquenta e três chegadas. Com treino nativo, só
o evolutivo o resolve, em quatro de vinte e uma execuções — abaixo do limiar
pré-registado. Os três mapas mostram um episódio de cada: o evolutivo atravessa
o labirinto até ao ninho; o PPO e o SAC ficam pelo caminho. A resposta é
negativa e está reportada como negativa. O que a composição degrada é a
fiabilidade: é a bimodalidade do Muro em U em ponto grande.
""", TERMOS["qi7"])

# 14. QI4 — mapa de escolha
slide_tabela(
    "QI4 — Critério de escolha", "não há vencedor universal; há um mapa de escolha",
    ["Se a missão…", "Então…", "Porquê"],
    [
        ["tem enxame de dimensão fixa e conhecida, e prioriza eficiência e cómputo",
         "MARL por gradiente (PPO)", "100 % em seis cenários, menor variância, ≈ 8× mais barato (12,8 vs 97,6 núcleos-hora)"],
        ["tem dimensão variável ou desconhecida, ou exige escalar sem retreino",
         "arquitetura de grafo com atenção (aqui, evolutiva)", "única a transferir de N=10 a N=100: 100 % em 28/28"],
        ["tem navegação estruturada (gargalos, salas, portas cooperativas)",
         "evolutivo com fitness de homing", "superior em 3 cenários, empate nos outros gargalos"],
        ["tem deceção espacial (gradiente enganador)",
         "neuroevolução + novidade doseada adaptativamente", "28/28 no Muro em U; sem custo nos outros seis"],
        ["exige tolerância a falhas de agentes", "qualquer dos três", "92–106 % de retenção — não discrimina"],
    ],
    larguras=[4.4, 3.6, 4.2],
    termos=TERMOS["qi4"],
    notas="""
A síntese prática. Dimensão fixa e cómputo barato: PPO. Dimensão variável ou
escala sem retreino: a arquitetura de grafo é a única opção entre as
estudadas. Navegação estruturada: o evolutivo com homing. Deceção espacial:
novidade doseada adaptativamente. Robustez a falhas não separa ninguém. Em vez
de um vencedor, a tese entrega os eixos que devem ditar a escolha.
""")

# 15. Limitações
slide_texto_figura(
    "Limitações — as que admito primeiro", "declaradas na dissertação, com os números",
    [
        "• Arquitetura assimétrica: atenção sobre grafo só no evolutivo; PPO/SAC com MLP — a comparação "
        "combina otimizador e representação (trabalho futuro n.º 1: a mesma arquitetura por gradiente)",
        "• Sete execuções por célula: onde a leitura depende de contagens, replicou-se a n = 28",
        "• Orçamento: 7 das 21 células ainda subiam no fim — o SAC nos gargalos lê-se como limite inferior "
        "(temperatura fixa, α = 0,1)",
        "• Só simulação: o fosso de implantação fica por validar",
        "• Dimensão vertical usada mas não observada; duas costuras na física das paredes — 11 de 343 "
        "modelos arquivados atravessam-nas, todos do evolutivo; na campanha final são 3 de 70 (Quatro "
        "Salas), e não as tornam melhores (55,7 vs 60,3 chegadas/ep). Onde a costura pesa é numa célula da "
        "campanha adaptativa — Quatro Salas —, declarada como contaminada",
        "• QI7 sobre um único mapa composto",
    ],
    figura="dotplot_eval_bottleneck.png", fig_w=5.2, fig_h=4.6, tamanho=15,
    termos=TERMOS["limitacoes"],
    notas="""
As limitações vêm antes das perguntas. A arquitetura difere entre paradigmas —
é a primeira, e é deliberadamente lida como o que permite isolar a
representação. Sete execuções são poucas onde a leitura é de contagens; por
isso o Muro em U foi replicado com vinte e oito. Sete células ainda subiam no
fim do orçamento; o SAC nos gargalos é limite inferior. Tudo é simulação. E a
física tem duas costuras nas paredes que uma minoria de modelos explora. Medi
todos os modelos arquivados, um a um: onze atravessam, todos do evolutivo, e na
campanha final são três, nas Quatro Salas — que não ficam melhores por isso, e
continuam muito acima do PPO nesse cenário. Onde a costura pesa mesmo é numa
célula da campanha de dosagem adaptativa, e essa está declarada como
contaminada, não usada como evidência.
""")

# 16. Contributos
slide_texto_figura(
    "Contributos", "",
    [
        ("1. O desenho da fitness como causa e cura do colapso evolutivo", {"negrito": True}),
        "  o homing terminal no potencial geodésico levou 28 execuções de 0 % a 100 % — o principal contributo metodológico",
        ("2. Caracterização condicional da procura por novidade", {"negrito": True}),
        "  resolve a deceção espacial (28/28) e, doseada adaptativamente, deixa de custar onde o gradiente basta",
        ("3. Um mapa de escolha para engenharia de enxames", {"negrito": True}),
        "  estrutura do cenário, variabilidade de N e cómputo, em vez de um vencedor universal",
        ("4. Uma implementação de referência aberta e verificável", {"negrito": True}),
        "  simulador, oito cenários, RS2C e neuroevolução; 27 verificadores ligam cada número aos dados",
    ],
    figura="escalabilidade_zeroshot_none.png", fig_w=5.4, fig_h=4.4,
    termos=TERMOS["contributos"],
    notas="""
Quatro contributos. O primeiro é metodológico: o colapso do evolutivo era o
sinal de treino, não o paradigma. O segundo é a caracterização da novidade como
instrumento direcionado — e a dosagem adaptativa que remove o seu custo. O
terceiro é o mapa de escolha. O quarto é o repositório: simulador, cenários,
controladores e os verificadores que ligam cada número da tese aos ficheiros de
avaliação.
""")

# 17. Conclusão
s = _novo_slide()
_titulo(s, "Conclusão", "a hipótese confirma-se só em parte — e não onde a punha")
_texto(s, MARGEM, Inches(1.9), W - 2 * MARGEM, Inches(4.5), [
    ("A vantagem de escala reside na representação — o grafo com atenção —, não no algoritmo de otimização.",
     {"tamanho": 22, "negrito": True, "cor": ACENTO}),
    ("Nos cenários de gradiente enganador, o que decide é o desenho do sinal de treino e a exploração "
     "doseada adaptativamente — não o paradigma.", {"tamanho": 22, "negrito": True, "cor": ACENTO}),
    ("Os métodos de gradiente ficam com a fiabilidade em espaço aberto e com ≈ 8× menos cómputo.",
     {"tamanho": 20}),
    ("A composição de dificuldades degrada a fiabilidade antes da magnitude — e isso está reportado "
     "como resultado negativo, com o número à vista.", {"tamanho": 20}),
], espaco=16)
_rodape(s)
_notas(s, """
A hipótese de que a inteligência adaptativa supera a robustez estática
confirma-se só em parte. A vantagem de escala existe — mas está na
representação, não no otimizador. Nos cenários enganadores, o que decide é o
sinal de treino e a exploração doseada. Os métodos de gradiente ficam com a
fiabilidade em espaço aberto e o cómputo barato. E a composição degrada a
fiabilidade antes da magnitude — um resultado negativo, reportado como tal.
""")

# 18. Demo ao vivo — o painel, no Muro em U
s = _novo_slide()
_titulo(s, "Demo ao vivo — o Muro em U no painel",
        "os três modelos a correr lado a lado, o vencedor em destaque, e a proveniência de cada número")
_fig(s, "demo_muro_em_u.jpg", MARGEM, Inches(1.7), Inches(8.3), Inches(4.7),
     pasta=os.path.join(RAIZ, "Defesa"))
_texto(s, MARGEM + Inches(8.6), Inches(1.7), W - 2 * MARGEM - Inches(8.6), Inches(4.9), [
    ("O que mostrar, por esta ordem", {"negrito": True, "cor": ACENTO, "tamanho": 15}),
    ("1. Os três da campanha final: 3/7, 4/7 e 2/7 — cada execução ou aprende o desvio ou fica a zero",
     {"tamanho": 13}),
    ("2. A 4.ª coluna: o GNN com novidade adaptativa — 7/7, 77,8 chegadas/ep (o «melhor»)", {"tamanho": 13}),
    ("3. Toggle «Episódio 3D»: os quatro enxames em 3D ao mesmo tempo; arrastar para rodar", {"tamanho": 13}),
    ("4. Em baixo: o dot plot com a linha do adaptativo, e as curvas de treino", {"tamanho": 13}),
    ("5. Se perguntarem «de onde vem este número?»: Proveniência — dois cliques até ao CSV",
     {"tamanho": 13}),
    ("", {"tamanho": 8}),
    ("Painel: localhost:8080 (portátil) · swarmroboticsgs.duckdns.org (Pi)", {"tamanho": 12, "cor": MUTED}),
    ("Setas do teclado mudam de mapa · plano B: a captura aqui ao lado, e a figura dos quatro braços "
     "do rodapé 10", {"tamanho": 12, "cor": MUTED}),
], tamanho=13, espaco=6)
_rodape(s)
_notas(s, """
Antes das perguntas, um minuto ao vivo. Abro o painel na Apresentação, mapa dois:
o Muro em U. Três colunas com os modelos da campanha final — repare-se que cada
execução ou aprende o desvio ou fica a zero: três em sete, quatro em sete, duas
em sete. A quarta coluna é o GNN com novidade doseada adaptativamente: sete em
sete, setenta e oito chegadas. Carrego em Episódio 3D e os quatro enxames
correm ao mesmo tempo. Em baixo, o dot plot com a linha do adaptativo e as
curvas de treino. Qualquer número do painel tem proveniência: dois cliques até
ao CSV. Plano B, se a rede ou o portátil falharem: a captura que está neste
slide, e a figura dos quatro braços — o slide com o rótulo 10 no rodapé.
""", TERMOS["demo"])

# 19. Fim
s = _novo_slide()
_barra = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGEM, Inches(2.5), Pt(3), Inches(0.85))
_barra.fill.solid()
_barra.fill.fore_color.rgb = INK_FORTE
_barra.line.fill.background()
_barra.shadow.inherit = False
_texto(s, MARGEM + Inches(0.32), Inches(2.4), W - 2 * MARGEM, Inches(1.2), "Obrigado.",
       tamanho=40, negrito=True, cor=INK_FORTE, fonte=TITULO_FT)
_texto(s, MARGEM, Inches(3.6), W - 2 * MARGEM, Inches(2.5), [
    ("Painel interativo (campanhas, vídeos, episódios 3D, proveniência de cada número):", {"cor": MUTED, "tamanho": 16}),
    ("http://swarmroboticsgs.duckdns.org", {"tamanho": 18, "cor": ACENTO}),
    ("Página de resultados:  https://pombo2001.github.io/swarm-robotics-tese/", {"tamanho": 16}),
    ("Código, dados e verificadores:  https://github.com/Pombo2001/swarm-robotics-tese", {"tamanho": 16}),
], espaco=8)
_notas(s, """
Obrigado. O painel tem as trinta campanhas, os vídeos, os episódios em 3D e a
proveniência de cada número — posso mostrar qualquer célula ao vivo. Estou à
disposição para perguntas.
""")

# A apresentação acaba aqui. O que vem a seguir não se mostra.
assert len(prs.slides) == len(PLANO), "o PLANO tem %d slides e a apresentação tem %d" % (len(PLANO), len(prs.slides))

# ── Slides de reserva ─────────────────────────────────────────────────────────
# Um a um, respondem às perguntas do docs/DEFESA_PERGUNTAS.md com a TABELA que o
# guião cita de memória. A diferença entre responder «cerca de oito vezes mais
# caro» e projetar os 97,6 contra os 12,8 é a diferença entre ter estudado e ter
# medido — e a pergunta pelo número exato é a que mais custa a improvisar em pé.
# Numeram-se A1…A6 e dizem-no no rodapé, para se poder pedir «o A2» em voz alta.
_ANEXO = True

# A1 — a tabela que a QI1 resume em três marcadores
slide_tabela(
    "Reserva · A tabela de avaliação, na íntegra",
    "chegadas/ep · taxa de sucesso · execuções a 100 %  —  7 execuções por célula, 20 episódios determinísticos cada",
    ["Cenário", "GNN evolutivo", "PPO", "SAC"],
    [
        ["Sandbox", "38,3 · 86 % · 5/7", "71,5 · 100 % · 7/7", "69,2 · 100 % · 7/7"],
        ["Muro em U", "24,5 · 43 % · 3/7", "39,6 · 71 % · 4/7", "9,0 · 34 % · 2/7"],
        ["Gargalo", "121,4 · 100 % · 7/7", "123,2 · 100 % · 7/7", "41,4 · 72 % · 5/7"],
        ["Quatro Salas", "59,8 · 100 % · 7/7", "33,6 · 100 % · 7/7", "31,8 · 100 % · 7/7"],
        ["Porta Cooperativa", "69,8 · 100 % · 7/7", "67,1 · 100 % · 7/7", "62,1 · 100 % · 7/7"],
        ["Perceção Cooperativa", "19,0 · 91 % · 6/7", "15,3 · 100 % · 7/7", "16,1 · 100 % · 7/7"],
        ["Porta com Alternativa", "86,7 · 100 % · 7/7", "85,3 · 100 % · 7/7", "68,6 · 100 % · 7/7"],
    ],
    larguras=[2.9, 3.1, 3.05, 3.05],
    bullets=[("Significância (Mann-Whitney sobre as médias por execução, n = 7; δ de Cliff): GNN > PPO e SAC "
              "nas Quatro Salas, Porta Cooperativa e Perceção (δ ≥ +0,71) · GNN > SAC no Gargalo e na Porta "
              "com Alternativa (δ = +1,00), empate com o PPO (p = 0,21) · PPO e SAC > GNN no Sandbox "
              "(δ = ±1,00) · Muro em U: nada significativo, nos três.", {"cor": MUTED, "tamanho": 14})],
    tamanho=12,
    notas="""
A tabela inteira, para quando a pergunta for por uma célula concreta. Quinze das
vinte e uma estão a cem por cento. As três parcelas de cada célula são chegadas
por episódio, taxa de sucesso e quantas das sete execuções chegaram aos cem por
cento — é a terceira que conta a história, porque é nela que a bimodalidade
aparece.
""")

# A2 — a QI6 braço a braço, incluindo a replicação a n = 28
slide_tabela(
    "Reserva · Novidade e mega-treino, braço a braço",
    "orçamento igualado a 195 min por execução, salvo indicação em contrário",
    ["Braço", "Muro em U (n = 7)", "Muro em U (n = 28)", "Porta c/ Alt. (n = 7)"],
    [
        ["GNN + novidade adaptativa", "7/7 · 68,5 ± 13,1", "28/28 · 67,4 ± 13,4", "77,2 ± 16,7  (n.s.)"],
        ["GNN + novidade w = 0,5 fixo", "7/7 · 69,8 ± 5,9", "—", "63,0 ± 21,9  (δ = −1,00)"],
        ["GNN objetivo puro", "3/7 · 24,5 ± 32,6", "15/28 · 32,5", "86,7 ± 2,0"],
        ["PPO", "4/7 · 39,6", "14/28 · 35,6", "85,3"],
        ["SAC", "2/7 · 9,0", "14/28 · 10,1", "68,6"],
    ],
    larguras=[3.5, 2.9, 2.95, 2.75],
    bullets=[
        "• A n = 28: Fisher exato, adaptativo vs. objetivo, p < 0,0001. GNN objetivo vs. PPO, p = 0,088 — "
        "indistinguíveis. Nenhuma das 28 execuções do SAC passa de 45,4 chegadas/ep: é uniformemente fraco, e "
        "não bimodal.",
        "• Não se compra com orçamento: o objetivo puro com 390 min continua bimodal (4/7, 31,5 ± 35,0); "
        "o adaptativo com 390 min faz 88,7 ± 0,6 na Porta com Alternativa — o melhor resultado da tese.",
        "• Ablação da temperagem, 4 variantes: 7/7 nos dois cenários — o mecanismo não vive da afinação.",
    ],
    tamanho=13,
    notas="""
Os quatro braços lado a lado. A coluna do meio é a replicação com vinte e oito
execuções por braço, que é onde uma leitura por contagens ganha poder. A última
coluna é o preço: com peso fixo, delta menos um; doseada adaptativamente, o
preço desaparece.
""")

# A3 — a QI7 em números, para a pergunta «um negativo vale o quê?»
slide_tabela(
    "Reserva · QI7, o mapa composto em números",
    "103 × 62 m · raio 60 m (contra 15 dos cenários isolados) · 128,8 m do spawn ao ninho · 106 obstáculos",
    ["Fase", "O que se mediu", "Resultado"],
    [
        ["Fase 1 — zero-shot", "transferência sem retreino, com 4 condições de controlo",
         "zero chegadas em 84/84 células · 1 680 episódios"],
        ["Controlo", "o mapa é sequer resolúvel?",
         "navegador geodésico, sem aprendizagem: 53,0 chegadas/ep (82,0 nas Quatro Salas)"],
        ["Fase 2 — treino nativo", "21 execuções por algoritmo, 780 min cada",
         "GNN 4/21 · 1,7 chegadas/ep  ·  PPO 0/21  ·  SAC 0/21"],
        ["Regra de decisão", "limiar fixado antes de haver dados",
         "15 execuções convergentes em 21 — o resultado ficou em 4"],
        ["Orçamento", "o treino tinha acabado de convergir?",
         "em 19 das 21 execuções a fitness ainda subia no último quinto"],
    ],
    larguras=[2.3, 4.5, 5.3],
    bullets=[("As 4 condições de controlo excluem a escala da observação, os obstáculos e as features da "
              "porta como causa do zero. O que a composição degrada é a fiabilidade, não a magnitude.",
              {"cor": MUTED, "tamanho": 14})],
    tamanho=12,
    notas="""
O negativo, com os números que o sustentam. A ordem importa: primeiro o zero em
oitenta e quatro células, depois a prova de que o mapa é resolúvel — cinquenta e
três chegadas de um navegador que não aprende nada —, e só então o limiar
pré-registado que não foi atingido. Sem o controlo do meio, o zero seria ambíguo.
""")

# A4 — o cómputo, para a pergunta «8× mais caro, é justo?»
slide_tabela(
    "Reserva · Custo, orçamento e o que está declarado",
    "a assimetria de cómputo, medida — e a leitura que dela se tira",
    ["", "GNN evolutivo", "PPO / SAC"],
    [
        ["Tempo por execução", "195 min", "48 min"],
        ["Paralelismo", "população de 30 genomas", "16 ambientes vetorizados"],
        ["Núcleos-hora por execução", "97,6", "12,8"],
        ["Rácio", "≈ 7,6× mais caro  (≈ 8×)", "—"],
    ],
    larguras=[3.9, 4.1, 4.1],
    bullets=[
        "• A leitura é a inversa da que favoreceria o evolutivo: com ≈ 8× menos núcleos-hora, os métodos "
        "de gradiente igualam ou superam o GNN em 4 dos 7 cenários.",
        "• Campanha principal: 147 treinos e 2 940 episódios de avaliação. Projeto inteiro: 28 sessões e "
        "2 078 h de treino, das quais 341 h do mega-treino de um mês.",
        "• Sub-treino declarado, não escondido: 7 das 21 células ainda subiam no fim do orçamento, e o SAC "
        "nos gargalos lê-se como limite inferior (temperatura fixa, α = 0,1, sem o ajuste dual — p. 44).",
    ],
    notas="""
Se a pergunta for pela justiça da comparação, é este o slide. A assimetria está
medida, e a conclusão que dela se tira é a que desfavorece o evolutivo: mesmo
custando oito vezes mais, não domina. E a configuração do SAC está declarada —
temperatura fixa, sem o ajuste dual —, pelo que os seus valores são um limite
inferior e não um veredicto sobre o algoritmo.
""")

# A5 — a prova visual da bimodalidade do Muro em U
s = _novo_slide()
_titulo(s, "Reserva · O Muro em U por dentro",
        "ocupação média do enxame ao longo do episódio — modelos campeões, 6 episódios cada")
_lbl = [("GNN evolutivo", GNN, "heatmap_ocupacao_gnn_u_wall.png"),
        ("PPO", PPO, "heatmap_ocupacao_ppo_u_wall.png"),
        ("SAC", SAC, "heatmap_ocupacao_sac_u_wall.png")]
_cw, _gap = Inches(3.95), Inches(0.14)
for _k, (_nome, _cor, _fich) in enumerate(_lbl):
    _x = MARGEM + _k * (_cw + _gap)
    _texto(s, _x, Inches(1.68), _cw, Inches(0.32), _nome, tamanho=15, negrito=True, cor=_cor,
           alinhar=PP_ALIGN.CENTER)
    _fig(s, _fich, _x, Inches(2.05), _cw, Inches(3.75))
_texto(s, MARGEM, H - Inches(1.2), W - 2 * MARGEM, Inches(0.6),
       "O corredor que contorna o muro é o que separa quem resolve de quem não resolve: o GNN e o PPO "
       "desenham-no (476 e 419 chegadas em 6 episódios); o SAC nunca o encontra e faz zero.",
       tamanho=13, cor=MUTED, alinhar=PP_ALIGN.CENTER)
_rodape(s)
_notas(s, """
Se perguntarem o que é, ao certo, resolver o Muro em U: é isto. O mapa de calor
mostra onde o enxame passa o episódio, no modelo campeão de cada algoritmo. O GNN
e o PPO desenham um corredor que sai do beco, contorna o muro e chega ao ninho —
quatrocentas e setenta e seis e quatrocentas e dezanove chegadas em seis episódios. O
SAC nunca encontra esse corredor: a ocupação espalha-se pela arena toda e o total
é zero. Não é uma diferença de magnitude, é outro comportamento — e é esta a
forma da falha que uma média esconde.
""")

# A6 — a pergunta «como sabemos que os números são os dos dados?»
slide_texto_figura(
    "Reserva · Como sei que os números são os dos dados", "a verificação, medida e ensaiada — não alegada",
    [
        "• 27 verificadores automáticos; 19 correm no hook de pre-commit e recusam o commit se um número "
        "da tese deixar de bater com o CSV que o produziu",
        "• O principal confere ~965 valores do main.tex; o do mapa composto, 72; o da configuração, 45 — "
        "física, recompensa e hiperparâmetros dos três algoritmos, lidos do foraging.yaml",
        "• Ensaio de mutação: a tese é estragada de propósito, 92 mutações, uma de cada vez, e o ensaio "
        "exige que todas sejam acusadas. Um verificador que nunca falhou não está provado — está por testar",
        "• Cobertura medida, não alegada: 1 043 dos 2 286 tokens numéricos do corpo do main.tex são lidos "
        "por algum verificador (46 %), e os 193 automatizáveis que faltam estão listados, um a um, em "
        "docs/COBERTURA_VERIFICADOR.md",
        "• Três pré-registos: hipótese, testes e regra de decisão escritos antes de haver dados",
        "• docs/REPRODUZIR.md refaz o percurso comando a comando — e é ele próprio ensaiado contra o disco",
    ],
    tamanho=16,
    notas="""
A resposta longa à pergunta da confiança. Não é que eu tenha conferido os
números: é que eles são recalculados a cada commit, e o commit é recusado se
algum deixar de bater. E os próprios verificadores são postos à prova — noventa
e duas mutações injetadas na tese, uma de cada vez, cada uma tem de ser
apanhada. A cobertura está medida e publicada, com a lista do que ainda falta,
porque um número de cobertura sem essa lista é uma alegação.
""")

prs.save(SAIDA)
print("[v] %s (%d slides: %d de apresentação + %d de reserva)"
      % (os.path.relpath(SAIDA, RAIZ), len(prs.slides), len(PLANO), _na))
print("    tempo planeado da apresentação: %s" % _mmss(sum(PLANO)))
